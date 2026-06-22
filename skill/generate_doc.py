#!/usr/bin/env python3
"""
Hackathon Technical Doc generator.

Produces a 1-page technical document (.docx) with a LOCKED structure and
styling so that every team's submission looks identical for the jury.

Teams supply ONLY content via a JSON spec. The layout, fonts, colors,
section order, and table shape are fixed by this script and cannot drift.

Usage:
    python generate_doc.py content.json [output.docx]

If output path is omitted, "<project_slug>-technical-doc.docx" is used.

See content.schema.json / example_content.json for the input format.
"""

import sys
import json
import re

try:
    from docx import Document
    from docx.shared import Pt, Emu, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    sys.stderr.write(
        "python-docx is required. Install with:\n"
        "    pip install python-docx --break-system-packages\n"
    )
    sys.exit(1)

# ---------------------------------------------------------------------------
# LOCKED DESIGN CONSTANTS  (extracted from the reference document)
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x1F, 0x38, 0x64)      # headings + title
GRAY = RGBColor(0x6D, 0x6D, 0x6D)      # subtitle + muted text
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "1F3864"                  # table header row fill
ALT_FILL = "EDF0F7"                     # zebra row fill

PAGE_W = Emu(7772400)                   # Letter width
PAGE_H = Emu(10058400)                  # Letter height
MARGIN = Emu(731520)                    # ~0.8in all sides

TITLE_PT = 17
SUBTITLE_PT = 9
HEADING_PT = 12.5
BODY_PT = 10
SMALL_PT = 8.5
FONT = "Calibri"

# Fixed, jury-facing labels. Structure is identical across languages.
LABELS = {
    "ka": {
        "doc_kind": "ტექნიკური დოკუმენტაცია",
        "s1": "მიმოხილვა",
        "s2": "არქიტექტურა",
        "s3": "ტექნოლოგიური სტეკი",
        "s4": "დანართი: საფეხურების დეტალები",
        "th": ["ფენა", "ტექნოლოგია", "რატომ ეს არჩევანი"],
        "principle": "პრინციპი",
        "appendix_note": "ქვემოთ მოცემული დეტალები არ წარმოადგენს დოკუმენტის ფოკუსს — იკითხება დამატებითი დროის არსებობისას.",
    },
    "en": {
        "doc_kind": "Technical Documentation",
        "s1": "Overview",
        "s2": "Architecture",
        "s3": "Technology Stack",
        "s4": "Appendix: Step Details",
        "th": ["Layer", "Technology", "Why this choice"],
        "principle": "Principle",
        "appendix_note": "The details below are not the focus of the document — read them only if time permits.",
    },
    "both": {
        "doc_kind": "ტექნიკური დოკუმენტაცია · Technical Documentation",
        "s1": "მიმოხილვა · Overview",
        "s2": "არქიტექტურა · Architecture",
        "s3": "ტექნოლოგიური სტეკი · Technology Stack",
        "s4": "დანართი: საფეხურები · Appendix: Steps",
        "th": ["ფენა · Layer", "ტექნოლოგია · Technology", "რატომ · Why"],
        "principle": "პრინციპი · Principle",
        "appendix_note": "ქვემოთ მოცემული დეტალები არ წარმოადგენს ფოკუსს · The details below are supplementary.",
    },
}


def _set_font(run, size=None, bold=None, color=None, name=FONT):
    run.font.name = name
    # Ensure Georgian glyphs use the same font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade_cell(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _no_space(paragraph, before=0, after=2, line=0.98):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def _heading(doc, text):
    p = doc.add_paragraph()
    _no_space(p, before=5, after=2)
    r = p.add_run(text)
    _set_font(r, size=HEADING_PT, bold=True, color=NAVY)
    return p


def _emit_rich(paragraph, text, base_size=BODY_PT, base_color=None):
    """Render **bold** spans inside a paragraph."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            _set_font(r, size=base_size, bold=True, color=base_color)
        else:
            r = paragraph.add_run(part)
            _set_font(r, size=base_size, bold=False, color=base_color)


def build(content, out_path):
    lang = content.get("lang", "ka")
    if lang not in LABELS:
        lang = "ka"
    L = LABELS[lang]

    doc = Document()
    # Page setup (locked)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = PAGE_W, PAGE_H
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = MARGIN

    # Default style
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)

    # --- Title block -------------------------------------------------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _no_space(p, before=0, after=1)
    r = p.add_run(content["title"])
    _set_font(r, size=TITLE_PT, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _no_space(p, before=0, after=6)
    subtitle = "  |  ".join(
        [s for s in [content.get("org", ""), content.get("event", ""), L["doc_kind"]] if s]
    )
    r = p.add_run(subtitle)
    _set_font(r, size=SUBTITLE_PT, bold=False, color=GRAY)

    # --- 1. Overview -------------------------------------------------------
    _heading(doc, f"1.  {L['s1']}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _no_space(p, after=3)
    _emit_rich(p, content["overview"])

    # --- 2. Architecture ---------------------------------------------------
    _heading(doc, f"2.  {L['s2']}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _no_space(p, after=3)
    flow = "  →  ".join(content["architecture_flow"])
    r = p.add_run(flow + ".")
    _set_font(r, size=BODY_PT)
    if content.get("architecture_principle"):
        r2 = p.add_run("  " + L["principle"] + ": ")
        _set_font(r2, size=BODY_PT, bold=True)
        r3 = p.add_run(content["architecture_principle"])
        _set_font(r3, size=BODY_PT)

    # --- 3. Technology Stack ----------------------------------------------
    _heading(doc, f"3.  {L['s3']}")
    if content.get("stack_summary"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _no_space(p, after=4)
        _emit_rich(p, content["stack_summary"])

    rows = content["stack_table"]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, htext in enumerate(L["th"]):
        _shade_cell(hdr[i], HEADER_FILL)
        cp = hdr[i].paragraphs[0]
        _no_space(cp, before=1, after=1)
        rr = cp.add_run(htext)
        _set_font(rr, size=SMALL_PT, bold=True, color=WHITE)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        if ri % 2 == 1:
            for c in cells:
                _shade_cell(c, ALT_FILL)
        for ci in range(3):
            cp = cells[ci].paragraphs[0]
            _no_space(cp, before=1, after=1)
            rr = cp.add_run(str(row[ci]))
            _set_font(rr, size=SMALL_PT, bold=(ci == 0))
            if ci == 0:
                rr.font.color.rgb = NAVY

    # --- 4. Appendix -------------------------------------------------------
    appendix = content.get("appendix", [])
    if appendix:
        _heading(doc, f"4.  {L['s4']}")
        p = doc.add_paragraph()
        _no_space(p, after=3)
        r = p.add_run(L["appendix_note"])
        _set_font(r, size=SMALL_PT, color=GRAY)
        for item in appendix:
            p = doc.add_paragraph(style="List Paragraph")
            _no_space(p, after=2)
            # item is "Label: description" or {"label":..,"text":..}
            if isinstance(item, dict):
                label, text = item.get("label", ""), item.get("text", "")
            else:
                if ":" in item:
                    label, text = item.split(":", 1)
                else:
                    label, text = "", item
            if label:
                r = p.add_run(label.strip() + ": ")
                _set_font(r, size=SMALL_PT, bold=True)
            r = p.add_run(text.strip())
            _set_font(r, size=SMALL_PT)

    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 2:
        sys.stderr.write("Usage: python generate_doc.py content.json [output.docx]\n")
        sys.exit(1)
    with open(sys.argv[1], encoding="utf-8") as f:
        content = json.load(f)
    if len(sys.argv) >= 3:
        out = sys.argv[2]
    else:
        slug = re.sub(r"[^a-zA-Z0-9]+", "-", content.get("title", "project")).strip("-").lower()
        out = f"{slug or 'project'}-technical-doc.docx"
    path = build(content, out)
    print(f"Saved: {path}")


if __name__ == "__main__":
    main()
